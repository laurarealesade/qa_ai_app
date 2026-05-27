"""
Processes the raw ServiceNow Excel and produces the aggregated dashboard table.

Input columns:  Created, Definition, ID, Value, Start, End, Duration, Calculation complete
  - ID    → case/ticket identifier
  - Value → agent ID

Output table columns: Agent_Name, Activity, Total_Toques, Total_Casos
  - Total_Toques → COUNT of rows per agent (every row = one toque)
  - Total_Casos  → DISTINCT COUNT of ID per agent
"""
import io
import logging
from pathlib import Path

import pandas as pd

from src.config import (
    AGENTS_REFERENCE_FILE,
    REF_ACTIVITY_COL,
    REF_AGENT_ID_COL,
    REF_AGENT_NAME_COL,
    SN_AGENT_ID_COL,
    SN_CASE_ID_COL,
)

log = logging.getLogger(__name__)


def _load_agents_reference() -> pd.DataFrame:
    if not AGENTS_REFERENCE_FILE.exists():
        raise FileNotFoundError(
            f"Archivo de referencia de agentes no encontrado: {AGENTS_REFERENCE_FILE}\n"
            "Crea data/agents_reference.xlsx con columnas: Agent_ID, Agent_Name, Activity"
        )
    df = pd.read_excel(AGENTS_REFERENCE_FILE, dtype={REF_AGENT_ID_COL: str})
    required = {REF_AGENT_ID_COL, REF_AGENT_NAME_COL, REF_ACTIVITY_COL}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"agents_reference.xlsx le faltan columnas: {missing}")
    return df[[REF_AGENT_ID_COL, REF_AGENT_NAME_COL, REF_ACTIVITY_COL]].drop_duplicates()


def process_servicenow_excel(file_bytes: bytes) -> pd.DataFrame:
    """
    Reads raw ServiceNow Excel bytes, merges with the agents reference table,
    and returns the aggregated dashboard DataFrame.
    """
    raw = pd.read_excel(io.BytesIO(file_bytes), dtype={SN_CASE_ID_COL: str, SN_AGENT_ID_COL: str})
    log.info("Raw rows loaded: %d", len(raw))

    # Normalize column names (strip whitespace)
    raw.columns = [c.strip() for c in raw.columns]

    if SN_CASE_ID_COL not in raw.columns or SN_AGENT_ID_COL not in raw.columns:
        raise ValueError(
            f"El Excel de ServiceNow debe tener las columnas '{SN_CASE_ID_COL}' y '{SN_AGENT_ID_COL}'. "
            f"Columnas encontradas: {list(raw.columns)}"
        )

    agents_ref = _load_agents_reference()

    merged = raw.merge(
        agents_ref,
        left_on=SN_AGENT_ID_COL,
        right_on=REF_AGENT_ID_COL,
        how="left",
    )

    unmatched = merged[REF_AGENT_NAME_COL].isna().sum()
    if unmatched:
        unknown_ids = merged.loc[merged[REF_AGENT_NAME_COL].isna(), SN_AGENT_ID_COL].unique()
        log.warning("%d filas con Agent_ID sin match en referencia: %s", unmatched, unknown_ids)
        merged[REF_AGENT_NAME_COL] = merged[REF_AGENT_NAME_COL].fillna("Unknown")
        merged[REF_ACTIVITY_COL] = merged[REF_ACTIVITY_COL].fillna("Unknown")

    aggregated = (
        merged.groupby([REF_AGENT_NAME_COL, REF_ACTIVITY_COL], as_index=False)
        .agg(
            Total_Toques=(SN_CASE_ID_COL, "count"),
            Total_Casos=(SN_CASE_ID_COL, "nunique"),
        )
        .rename(columns={REF_AGENT_NAME_COL: "Agent_Name", REF_ACTIVITY_COL: "Activity"})
        .sort_values(["Activity", "Agent_Name"])
        .reset_index(drop=True)
    )

    log.info("Aggregated rows: %d agents", len(aggregated))
    return aggregated


def dataframe_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Serializes the dashboard DataFrame to Excel bytes with basic formatting."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Dashboard")
        ws = writer.sheets["Dashboard"]

        # Column widths
        col_widths = {"A": 30, "B": 15, "C": 18, "D": 16}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        # Header style
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill(fill_type="solid", fgColor="1F3864")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Activity color bands
        activity_colors = {"Online": "D9EAD3", "Offline": "FCE5CD", "7Eleven": "CFE2F3"}
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            activity = row[1].value  # column B
            color = activity_colors.get(activity, "FFFFFF")
            fill = PatternFill(fill_type="solid", fgColor=color)
            for cell in row:
                cell.fill = fill

    return buf.getvalue()
